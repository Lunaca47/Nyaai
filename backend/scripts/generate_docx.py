from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from bs4 import BeautifulSoup
import os

def create_docx(html_path, output_path, kiosk_img, workflow_img):
    doc = Document()
    
    # Define styles
    style = doc.styles['Normal']
    font = style.font
    font.name = 'Calibri'
    font.size = Pt(11)

    # 1. Cover Page
    with open(html_path, 'r', encoding='utf-8') as f:
        soup = BeautifulSoup(f, 'html.parser')

    # Title
    title = soup.find(class_='cover-title').get_text(strip=True)
    subtitle = soup.find(class_='cover-sub').get_text(strip=True)
    
    doc.add_paragraph('\n' * 5)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(title)
    run.bold = True
    run.font.size = Pt(36)
    run.font.color.rgb = RGBColor(0, 0, 0) # Black

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(subtitle)
    run.font.size = Pt(16)
    run.italic = True
    run.font.color.rgb = RGBColor(0, 0, 0) # Black

    doc.add_paragraph('\n' * 2)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run("Patent Concept Document · Draft v1.0 (Monochrome)")
    run.font.size = Pt(10)
    run.font.color.rgb = RGBColor(0, 0, 0) # Black

    doc.add_page_break()

    # 2. Content Sections
    sections = soup.find_all(class_='section')
    for section in sections:
        section_id = section.get('id', '')
        header = section.find(class_='section-header')
        if header:
            num = header.find(class_='section-num').get_text(strip=True)
            title_text = header.find(class_='section-title').get_text(strip=True)
            
            # Heading
            h = doc.add_heading(f"{num}. {title_text}", level=1)
            # Custom style for heading
            for run in h.runs:
                run.font.color.rgb = RGBColor(0, 0, 0) # Black
        
        # Highlight box
        highlight = section.find(class_='highlight-box')
        if highlight:
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Inches(0.2)
            run = p.add_run(highlight.get_text(strip=True))
            run.italic = True
            run.font.color.rgb = RGBColor(0, 0, 0) # Black

        # Prose paragraphs
        prose = section.find(class_='prose')
        if prose:
            for p_html in prose.find_all('p'):
                para = doc.add_paragraph(p_html.get_text(strip=True))
                para.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY

        # Handle Diagrams
        if section_id == 's3' and kiosk_img:
            doc.add_picture(kiosk_img, width=Inches(4))
            last_p = doc.paragraphs[-1]
            last_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            doc.add_paragraph("Figure 1: NYAAI Kiosk Physical Layout (Monochrome)", style='Caption').alignment = WD_ALIGN_PARAGRAPH.CENTER

        if section_id == 's5' and workflow_img:
            doc.add_picture(workflow_img, width=Inches(4.5))
            last_p = doc.paragraphs[-1]
            last_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            doc.add_paragraph("Figure 2: User Workflow Flowchart (Monochrome)", style='Caption').alignment = WD_ALIGN_PARAGRAPH.CENTER

        # Lists handled simply
        workflow_steps = section.find(class_='workflow-steps')
        if workflow_steps:
            for step in workflow_steps.find_all(class_='step'):
                step_text = step.find(class_='step-text')
                title = step_text.find('strong').get_text(strip=True)
                desc = step_text.find('span').get_text(strip=True)
                doc.add_paragraph(f"• {title}: {desc}", style='List Bullet')

        # Cards and Grids
        cards = section.find_all(class_='card')
        if cards:
            table = doc.add_table(rows=0, cols=2)
            table.style = 'Table Grid'
            for card in cards:
                label = card.find(class_='card-label')
                value = card.find(class_='card-value')
                if label and value:
                    row_cells = table.add_row().cells
                    row_cells[0].text = label.get_text(strip=True)
                    row_cells[1].text = value.get_text(strip=True)

    doc.save(output_path)
    print(f"Document saved to {output_path}")

if __name__ == "__main__":
    HTML_FILE = r"c:\Users\barma\.gemini\antigravity\scratch\nyaai_app\nyaai_concept.html"
    OUTPUT_FILE = r"c:\Users\barma\.gemini\antigravity\scratch\nyaai_app\NYAAI_Patent_Concept_BW.docx"
    KIOSK_IMG = r"C:\Users\barma\.gemini\antigravity\brain\546af03b-1f58-4686-8c92-ca3caba165e4\kiosk_layout_diagram_bw_png_1775819410114.png"
    WORKFLOW_IMG = r"C:\Users\barma\.gemini\antigravity\brain\546af03b-1f58-4686-8c92-ca3caba165e4\user_workflow_flowchart_bw_png_1775819443787.png"
    
    create_docx(HTML_FILE, OUTPUT_FILE, KIOSK_IMG, WORKFLOW_IMG)
